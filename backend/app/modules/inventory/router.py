from datetime import date, timedelta
from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import get_current_user
from .model import InventoryItem, InventoryMovement

router = APIRouter(prefix="/api/inventory", tags=["Inventory"])


def require_admin(user):
    if getattr(getattr(user, "role", None), "name", None) != "Administrador":
        raise HTTPException(status_code=403, detail="Solo el administrador puede gestionar el inventario")


class ItemIn(BaseModel):
    name: str = Field(min_length=1, max_length=160)
    sku: str | None = None
    barcode: str | None = None
    category: str = "General"
    item_type: str = "INGREDIENTE"
    brand: str | None = None
    presentation: str | None = None
    unit: str = "unidad"
    quantity: float = Field(default=0, ge=0)
    min_quantity: float = Field(default=0, ge=0)
    max_quantity: float = Field(default=0, ge=0)
    reorder_quantity: float = Field(default=0, ge=0)
    unit_cost: float = Field(default=0, ge=0)
    storage: str | None = None
    location: str | None = None
    purchase_date: date | None = None
    opened_date: date | None = None
    expiry_date: date | None = None
    storage_temperature: str | None = None
    supplier: str | None = None
    allergen: str | None = None
    notes: str | None = None


class MovementIn(BaseModel):
    movement_type: str
    quantity: float = Field(gt=0)
    reason: str | None = None


def item_status(item: InventoryItem):
    qty = float(item.quantity or 0)
    minimum = float(item.min_quantity or 0)
    today = date.today()
    if qty <= 0:
        return "SIN STOCK"
    if item.expiry_date and item.expiry_date < today:
        return "VENCIDO"
    if item.expiry_date and item.expiry_date <= today + timedelta(days=7):
        return "REVISAR VENCIMIENTO"
    if minimum > 0 and qty <= minimum:
        return "POR AGOTARSE"
    return "OK"


def recommendation(item: InventoryItem, status: str):
    if status == "VENCIDO":
        return "Retirar de servicio y revisar cadena de frío."
    if status == "REVISAR VENCIMIENTO":
        return "Revisar fecha y cadena de frío; priorizar consumo si corresponde."
    if status == "SIN STOCK":
        return "Reponer antes del próximo servicio."
    if status == "POR AGOTARSE":
        qty = float(item.reorder_quantity or 0)
        return f"Programar reposición{f' de {qty:g} {item.unit}' if qty else ''}."
    return "Sin alerta."


def serialize_item(item: InventoryItem):
    status = item_status(item)
    return {
        "id": str(item.id),
        "name": item.name,
        "sku": item.sku,
        "barcode": item.barcode,
        "category": item.category,
        "item_type": item.item_type,
        "brand": item.brand,
        "presentation": item.presentation,
        "unit": item.unit,
        "quantity": float(item.quantity or 0),
        "min_quantity": float(item.min_quantity or 0),
        "max_quantity": float(item.max_quantity or 0),
        "reorder_quantity": float(item.reorder_quantity or 0),
        "unit_cost": float(item.unit_cost or 0),
        "storage": item.storage,
        "location": item.location,
        "purchase_date": item.purchase_date.isoformat() if item.purchase_date else None,
        "opened_date": item.opened_date.isoformat() if item.opened_date else None,
        "expiry_date": item.expiry_date.isoformat() if item.expiry_date else None,
        "storage_temperature": item.storage_temperature,
        "supplier": item.supplier,
        "allergen": item.allergen,
        "notes": item.notes,
        "status": status,
        "recommendation": recommendation(item, status),
        "value": round(float(item.quantity or 0) * float(item.unit_cost or 0), 2),
    }


@router.get("")
def list_items(db: Session = Depends(get_db), user=Depends(get_current_user)):
    require_admin(user)
    items = db.query(InventoryItem).filter(InventoryItem.active.is_(True)).order_by(InventoryItem.category, InventoryItem.name).all()
    data = [serialize_item(i) for i in items]
    counts = {status: sum(x["status"] == status for x in data) for status in ("OK", "POR AGOTARSE", "SIN STOCK", "REVISAR VENCIMIENTO", "VENCIDO")}
    return {
        "items": data,
        "total": len(data),
        "counts": counts,
        "low_stock": counts["SIN STOCK"] + counts["POR AGOTARSE"],
        "expiry_alerts": counts["VENCIDO"] + counts["REVISAR VENCIMIENTO"],
        "stock_value": round(sum(x["value"] for x in data), 2),
    }


@router.post("")
def create_item(body: ItemIn, db: Session = Depends(get_db), user=Depends(get_current_user)):
    require_admin(user)
    data = body.model_dump()
    for field in ("sku", "barcode"):
        if data[field]:
            exists = db.query(InventoryItem).filter(getattr(InventoryItem, field) == data[field]).first()
            if exists:
                raise HTTPException(409, f"El {field.upper()} ya existe en el inventario.")
        else:
            data[field] = None
    if data["max_quantity"] and data["max_quantity"] < data["min_quantity"]:
        raise HTTPException(400, "El stock máximo no puede ser menor que el mínimo.")
    item = InventoryItem(**data)
    db.add(item)
    db.flush()
    if item.quantity:
        db.add(InventoryMovement(item_id=item.id, movement_type="ENTRADA", quantity=item.quantity, reason="Inventario inicial"))
    db.commit()
    db.refresh(item)
    return serialize_item(item)


@router.get("/report.xlsx")
def report(db: Session = Depends(get_db), user=Depends(get_current_user)):
    require_admin(user)
    items = db.query(InventoryItem).filter(InventoryItem.active.is_(True)).order_by(InventoryItem.category, InventoryItem.name).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    headers = [
        "Producto", "SKU", "Código de barras", "Categoría", "Tipo", "Marca", "Presentación", "Unidad",
        "Existencia", "Stock mínimo", "Stock máximo", "Cantidad a reponer", "Costo unitario", "Valor existencia",
        "Almacenamiento", "Ubicación", "Compra", "Apertura", "Vencimiento", "Temperatura",
        "Proveedor", "Alérgenos", "Estado", "Recomendación", "Notas",
    ]
    ws.append(headers)
    for item in items:
        p = serialize_item(item)
        ws.append([
            item.name, item.sku, item.barcode, item.category, item.item_type, item.brand, item.presentation, item.unit,
            p["quantity"], p["min_quantity"], p["max_quantity"], p["reorder_quantity"], p["unit_cost"], p["value"],
            item.storage, item.location, item.purchase_date, item.opened_date, item.expiry_date,
            item.storage_temperature, item.supplier, item.allergen, p["status"], p["recommendation"], item.notes,
        ])
    style_sheet(ws, "202A44")

    alerts = wb.create_sheet("Alertas")
    alerts.append(["Producto", "Existencia", "Mínimo", "Almacenamiento", "Ubicación", "Vencimiento", "Estado", "Recomendación"])
    for item in items:
        p = serialize_item(item)
        if p["status"] != "OK":
            alerts.append([item.name, p["quantity"], p["min_quantity"], item.storage, item.location, item.expiry_date, p["status"], p["recommendation"]])
    style_sheet(alerts, "9B1C1C")

    movements = wb.create_sheet("Movimientos")
    movements.append(["Fecha", "Producto", "Tipo", "Cantidad", "Unidad", "Motivo"])
    movement_rows = (
        db.query(InventoryMovement, InventoryItem)
        .join(InventoryItem, InventoryItem.id == InventoryMovement.item_id)
        .order_by(InventoryMovement.created_at.desc())
        .limit(5000)
        .all()
    )
    for movement, item in movement_rows:
        movements.append([movement.created_at, item.name, movement.movement_type, float(movement.quantity or 0), item.unit, movement.reason])
    style_sheet(movements, "3B5B8A")

    output = BytesIO()
    wb.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=informe_inventario.xlsx"},
    )



@router.put("/{item_id}")
def update_item(item_id: UUID, body: ItemIn, db: Session = Depends(get_db), user=Depends(get_current_user)):
    require_admin(user)
    item = db.query(InventoryItem).filter(InventoryItem.id == item_id, InventoryItem.active.is_(True)).first()
    if not item:
        raise HTTPException(404, "Producto de inventario no encontrado")
    data = body.model_dump()
    if data["max_quantity"] and data["max_quantity"] < data["min_quantity"]:
        raise HTTPException(400, "El stock máximo no puede ser menor que el mínimo.")
    for field in ("sku", "barcode"):
        value = data[field] or None
        if value:
            exists = db.query(InventoryItem).filter(getattr(InventoryItem, field) == value, InventoryItem.id != item.id).first()
            if exists:
                raise HTTPException(409, f"El {field.upper()} ya existe en el inventario.")
        data[field] = value
    for key, value in data.items():
        setattr(item, key, value)
    db.commit()
    db.refresh(item)
    return serialize_item(item)


@router.post("/{item_id}/movement")
def register_movement(item_id: UUID, body: MovementIn, db: Session = Depends(get_db), user=Depends(get_current_user)):
    require_admin(user)
    item = db.query(InventoryItem).filter(InventoryItem.id == item_id, InventoryItem.active.is_(True)).first()
    if not item:
        raise HTTPException(404, "Producto no encontrado")
    movement_type = body.movement_type.upper()
    if movement_type not in ("ENTRADA", "SALIDA", "MERMA", "AJUSTE"):
        raise HTTPException(400, "Tipo de movimiento inválido")
    qty = float(body.quantity)
    if movement_type == "AJUSTE":
        new_quantity = qty
    else:
        new_quantity = float(item.quantity or 0) + (qty if movement_type == "ENTRADA" else -qty)
    if new_quantity < 0:
        raise HTTPException(400, "No hay existencias suficientes para esa salida")
    item.quantity = new_quantity
    db.add(InventoryMovement(item_id=item.id, movement_type=movement_type, quantity=qty, reason=body.reason))
    db.commit()
    db.refresh(item)
    return serialize_item(item)


@router.get("/{item_id}/movements")
def item_movements(item_id: UUID, db: Session = Depends(get_db), user=Depends(get_current_user)):
    require_admin(user)
    item = db.query(InventoryItem).filter(InventoryItem.id == item_id).first()
    if not item:
        raise HTTPException(404, "Producto no encontrado")
    rows = db.query(InventoryMovement).filter(InventoryMovement.item_id == item_id).order_by(InventoryMovement.created_at.desc()).limit(100).all()
    return {"item": serialize_item(item), "movements": [{"id": str(m.id), "type": m.movement_type, "quantity": float(m.quantity or 0), "reason": m.reason, "created_at": m.created_at} for m in rows]}


@router.delete("/{item_id}")
def archive_item(item_id: UUID, db: Session = Depends(get_db), user=Depends(get_current_user)):
    require_admin(user)
    item = db.query(InventoryItem).filter(InventoryItem.id == item_id).first()
    if not item:
        raise HTTPException(404, "Producto no encontrado")
    item.active = False
    db.commit()
    return {"ok": True}



def style_sheet(ws, color):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in ws.columns:
        letter = column[0].column_letter
        width = max((len(str(cell.value or "")) for cell in column), default=10) + 2
        ws.column_dimensions[letter].width = min(max(width, 12), 34)
